from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(r"Docs/ML/Briefing_Mentoria_Tecnica_ClimaGrid.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.75)
sec.right_margin = Inches(0.75)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(10.2)
styles['Normal'].paragraph_format.space_after = Pt(5)
styles['Normal'].paragraph_format.line_spacing = 1.08
for name, size in [('Title', 23), ('Heading 1', 15), ('Heading 2', 11.5), ('Heading 3', 10.5)]:
    st = styles[name]
    st.font.name = 'Aptos Display' if name != 'Heading 3' else 'Aptos'
    st._element.rPr.rFonts.set(qn('w:ascii'), st.font.name)
    st._element.rPr.rFonts.set(qn('w:hAnsi'), st.font.name)
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(10 if name == 'Heading 1' else 6)
    st.paragraph_format.space_after = Pt(4)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def borders(table, color='D9D9D9'):
    tblPr = table._tbl.tblPr
    b = tblPr.first_child_found_in('w:tblBorders')
    if b is None:
        b = OxmlElement('w:tblBorders')
        tblPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge
        el = b.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); b.append(el)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:space'), '0'); el.set(qn('w:color'), color)

def cell_text(cell, text, bold=False, size=8.5, color='000000'):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(str(text)); r.bold = bold; r.font.name = 'Aptos'; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(headers, rows, widths=None, font=8.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    borders(t)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    trPr.append(tbl_header)
    for i, h in enumerate(headers):
        cell_text(t.rows[0].cells[i], h, True, font, 'FFFFFF'); shade(t.rows[0].cells[i], '263B68')
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cell_text(cells[i], value, False, font)
            if ri % 2 == 1: shade(cells[i], 'F3F5F8')
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def p(text='', bold_lead=None):
    para = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = para.add_run(bold_lead); r.bold = True
        para.add_run(text[len(bold_lead):])
    else:
        para.add_run(text)
    return para

def bullet(text, level=0):
    para = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    para.paragraph_format.space_after = Pt(2)
    para.add_run(text)
    return para

def numbered(text):
    para = doc.add_paragraph(style='List Number')
    para.paragraph_format.space_after = Pt(2)
    para.add_run(text)
    return para

def page_break(): doc.add_page_break()

# Footer
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run('ClimaGrid | Briefing para mentoria técnica | ')
fr.font.size = Pt(8); fr.font.color.rgb = RGBColor(100,100,100)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
footer._p.append(fld)

# Title block
title = doc.add_paragraph(style='Title')
title.add_run('ClimaGrid e o estimador vento para potência')
sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(12)
r = sub.add_run('Briefing de mentoria técnica sobre escopo, estado do ML e viabilidade de uso'); r.italic = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor(70,70,70)
p('Objetivo. Preparar uma conversa técnica sobre o que o ClimaGrid já implementou, o que ainda é experimento e quais decisões científicas precisam ser aprovadas antes de apresentar o modelo como aplicável.')
p('Conclusão em uma frase. O projeto já demonstra uma cadeia útil de cenário climático até um PWF de uma hora, mas o ML ainda não substitui a curva física em produção: hoje a curva é o estimador servido na fase 2, enquanto o LightGBM híbrido é um experimento que precisa de alvo aprovado, histórico mais amplo, validação temporal intacta e confirmação no ANAREDE.')

doc.add_heading('1 O que o ClimaGrid resolve', level=1)
p('O ClimaGrid recebe uma condição de vento associada aos conjuntos eólicos do subsistema Nordeste, estima uma injeção de potência ativa em MW, distribui essa potência para as barras elétricas mapeadas e exporta um PWF para análise no ANAREDE. O produto prepara o caso; não executa o fluxo de potência nem afirma convergência elétrica.')
p('A unidade operacional é uma hora e um PWF representa um único instante. Um período gera uma coleção de instantes e arquivos, não um PWF temporal único. No replay histórico, a geração usada é a geração observada da ONS; esse fluxo não depende de ML.')

doc.add_heading('2 Evolução do projeto e fases', level=1)
table(['Fase','Estado verificado','Entrada e saída','Papel do ML'], [
    ['1 Replay histórico','Implementada','ONS GERACAO_USINA-2_HO + ERA5 do mesmo instante; exporta PWF com Pg observado.','Nenhum. Não chamar observação de estimativa.'],
    ['2 Cenário climático','Protótipo ponta a ponta em validação; limite do MVP','CSV normalizado ou uma hora histórica ERA5; estimativa por conjunto; revisão de barras; PWF de uma hora.','Curva física servida hoje. O híbrido LightGBM ainda não está integrado ao caminho do cenário.'],
    ['3 Hora futura','Pós-MVP','Exige fonte explícita de meteorologia futura, climatologia ou cenário probabilístico.','Só depois da fase 2; ERA5 isolado não é previsão futura.'],
    ['4 Curtailment','Fora do MVP; por último','Separaria potencial, geração observada, indisponibilidade e causa de restrição.','Não deve ser antecipado nem confundido com estimativa bruta.'],
], [0.95,1.35,3.05,1.85], 8.2)
p('Os PDFs de visão do Ideathon descrevem uma ambição maior, incluindo previsão futura, risco de curtailment e classificador de causa. O contexto canônico atual restringe o MVP à fase 2. Para a mentoria, a leitura correta é usar os PDFs como origem da proposta e o repositório atual como estado de implementação.')

doc.add_heading('3 O que está implementado hoje', level=1)
p('Fato verificado no código. O caminho de upload e de ERA5 histórico identifica o estimador como physical-curve-v1 em backend/ai-service/app/climate_file.py. Ele calcula a curva física com velocidade do vento, capacidade e disponibilidade, aplica cut-in/cut-out e limita o resultado a capacidade instalada multiplicada pela disponibilidade.')
p('Fato verificado no código. Existe um Predictor em backend/ai-service/app/predictor.py e um treino em backend/ai-service/training/train.py. O treino usa LightGBM para aprender um resíduo em relação à curva física, em fração da capacidade. O artefato só é considerado aprovado quando o MAE do híbrido no teste é menor que o MAE da curva física; se o artefato faltar, estiver inválido, estiver fora do domínio ou não estiver aprovado, a API cai para a curva física.')
p('Limite operacional. Essa regra de aprovação é uma regra experimental de artefato, não uma aprovação científica para produção. O endpoint da fase 2 ainda não deve ser descrito como “IA em operação”.')

doc.add_heading('4 Como o modelo preditivo agiria', level=1)
p('O desenho atual é híbrido e residual:')
numbered('O cenário fornece uma hora, o conjunto ONS e u100/v100 em m/s; temperatura, pressão e disponibilidade podem completar o contexto conforme o contrato.')
numbered('A curva física produz o baseline em MW, com zero abaixo do cut-in, rampa até a potência nominal, patamar e zero acima do cut-out.')
numbered('O LightGBM estima uma correção do baseline usando variáveis climáticas, calendário, capacidade, disponibilidade e indicadores de qualidade. O modelo não deveria usar uma restrição contemporânea como feature do cenário.')
numbered('A correção é somada ao baseline e novamente limitada a [0, capacidade x disponibilidade]. Essa barreira é a garantia física disponível no código.')
numbered('A estimativa por conjunto é distribuída por CEG para uma ou mais barras, e o writer altera somente Pg no PWF. A validação elétrica continua no ANAREDE.')
p('Interpretação importante. Como a saída é limitada por capacidade x disponibilidade, a grandeza atual se comporta como potência disponível condicionada à disponibilidade informada. Isso não é automaticamente potencial bruto eólico, nem geração verificada observada.')

doc.add_heading('5 Quatro grandezas que não podem ser misturadas', level=1)
table(['Grandeza','O que representa','Fonte ou hipótese','Uso adequado'], [
    ['Potencial eólico bruto','Produção associada ao vento, sem aplicar indisponibilidade operacional ou corte de rede.','Curva física/modelo; requer disponibilidade igual a 1 ou outro contrato explícito.','Estimar o recurso e separar efeitos operacionais.'],
    ['Potência disponível','O que poderia ser gerado no instante dado o vento e a disponibilidade eletromecânica.','Vento + cadastro + disponibilidade válida.','Alvo mais compatível com o clipping atual, se o ONS confirmar a semântica.'],
    ['Geração efetivamente verificada','Produção observada/entregue, podendo conter efeitos de manutenção, restrição e operação.','ONS GERACAO_USINA-2_HO.','Replay e análise operacional; não é automaticamente potencial.'],
    ['Geração de referência ONS','Estimativa ONS usada para representar geração sem limitação, segundo a documentação do piloto.','RESTRICAO_COFF_EOLICA e coluna geracao_referencia_mw.','Candidato a proxy de potencial, mas exige decisão sobre excessos e qualidade.'],
], [1.25,2.25,2.15,1.55], 8.0)
p('Decisão pendente. geracao_referencia_mw e geracao_verificada_mw estão ambos aceitos pelo pipeline de treino, mas não são alvos equivalentes. A referência ONS não é uma medição direta e pode superar capacidade instalada ou capacidade disponível. A geração verificada é observada, porém mistura vento com restrições e indisponibilidades. O especialista ONS precisa aprovar o estimando, os filtros e a política de elegibilidade antes de qualquer promoção.')

doc.add_heading('6 O que os experimentos realmente mostram', level=1)
p('Os resultados abaixo são evidências de viabilidade do pipeline, não validação de produção. O piloto multiusina usa agosto de 2024, alvo geracao_referencia_mw e um split temporal. O teste já foi consultado em várias rodadas e não deve orientar novas escolhas sem reservar outro período final.')
table(['Experimento','Validação física','Validação híbrida','Teste físico','Teste híbrido','Leitura'], [
    ['Usina 01 exp 1','n/d','n/d','73,77 MW','72,30 MW','Melhora pequena; aprovado apenas pela regra experimental.'],
    ['Usina 01 exp 2','n/d','n/d','73,77 MW','80,10 MW','Piorou; artefato rejeitado e fallback preservado.'],
    ['Multiusina 03','49,77 MW','20,78 MW','51,03 MW','23,63 MW','Melhor teste entre 03-06; ainda um mês.'],
    ['Multiusina 04','49,77 MW','20,94 MW','51,03 MW','24,03 MW','Melhora; não venceu o teste.'],
    ['Multiusina 05','49,77 MW','20,77 MW','51,03 MW','23,71 MW','Melhor validação por 0,006 MW; não prova superioridade.'],
    ['Multiusina 06','49,77 MW','20,89 MW','51,03 MW','23,98 MW','Melhora; ainda dependente do mesmo piloto.'],
], [1.05,1.0,1.05,0.9,0.95,2.05], 7.6)
p('O intervalo empírico do experimento multiusina cobriu cerca de 85% do teste, abaixo dos 90% nominais. Há também grande variação por conjunto. Isso reforça a necessidade de medir por mês, usina, faixa de vento e disponibilidade, em vez de aceitar somente o MAE agregado.')

doc.add_heading('7 A pergunta central sobre substituir a curva física', level=1)
doc.add_heading('Resposta curta', level=2)
p('O modelo pode se tornar aplicável como estimador de apoio à decisão, mas ainda não há evidência para substituir a curva física hoje. A recomendação técnica é manter a curva como baseline, limite e fallback, e permitir que o ML a corrija somente quando um artefato aprovado estiver dentro do domínio de treino.')
doc.add_heading('Por que o ML é interessante', level=2)
bullet('A curva genérica assume uma relação idealizada entre vento e potência. Ela não conhece bem composição da frota, esteira, topografia, densidade do ar, viés do ponto ERA5, degradação ou padrões locais.')
bullet('O modelo residual pode aprender apenas a parte sistemática que a curva não explica, reduzindo viés sem abandonar a restrição física.')
bullet('Um modelo aprovado pode produzir intervalos de incerteza e sinalizar cenários fora do domínio, algo que uma curva fixa não oferece.')
bullet('A melhora só importa se alterar decisões de estudo: por exemplo, a injeção estimada precisa mudar a conclusão do fluxo de potência ou reduzir retrabalho do analista.')
doc.add_heading('Por que não remover a curva', level=2)
bullet('O ML depende de um alvo com semântica correta e de histórico representativo. Um único mês não cobre sazonalidade, mudanças cadastrais e regimes de vento.')
bullet('A geração verificada inclui efeitos que não são vento; treinar diretamente nela pode ensinar restrição ao estimador de potencial.')
bullet('A referência ONS é útil, mas não é verdade absoluta. Linhas acima de capacidade disponível precisam de classificação ou revisão, não de clipping silencioso.')
bullet('A curva oferece comportamento seguro fora da distribuição e uma explicação simples para engenharia. Retirá-la elimina uma barreira útil antes de provar generalização e aceitação no ANAREDE.')

doc.add_heading('8 Condições para uma eventual substituição', level=1)
numbered('Aprovar o contrato do alvo: potencial bruto, potência disponível ou geração verificada; unidade MW, uma hora, conjunto ONS e política de disponibilidade.')
numbered('Construir snapshot reproduzível com vários meses, ciclo anual completo e, se possível, anos adicionais, mantendo a origem por CEG, vigência de capacidade/composição e UTC.')
numbered('Comparar curva física, híbrido residual e alternativas diretas com splits temporais bloqueados, um período final intocado e avaliação por conjunto e regime.')
numbered('Demonstrar ganho estável em MAE/RMSE e erro relativo, mas também cobertura de intervalos, limites físicos, out-of-domain e impacto no PWF.')
numbered('Validar um conjunto de cenários no ANAREDE. O critério não é afirmar que o ML converge; é confirmar que o PWF preserva a estrutura e que a análise elétrica continua válida.')
numbered('Operar primeiro em shadow mode, registrando saída da curva e do ML sem alterar a decisão do usuário. Só depois de monitoramento e aprovação de domínio promover o artefato.')

doc.add_heading('9 Perguntas para o mentor técnico', level=1)
doc.add_heading('Alvo e domínio', level=2)
bullet('Qual é o estimando de negócio: potencial bruto, potência disponível ou geração verificada? O valor esperado deve ser limitado por capacidade ou por capacidade x disponibilidade?')
bullet('A coluna geracao_referencia_mw pode ser tratada como proxy de produção sem limitação? O que explica linhas acima da capacidade instalada e acima da capacidade disponível?')
bullet('Existe uma versão oficial, regra de cálculo ou campo de qualidade para a referência ONS? Quais registros devem ser excluídos, auditados ou mantidos com flag?')
bullet('A disponibilidade significa capacidade eletromecânica disponível, disponibilidade declarada no ONS ou outra medida? Ela é conhecida no momento de uso do cenário?')
doc.add_heading('Dados e generalização', level=2)
bullet('Qual composição de conjuntos ONS e quais CEGs devem ser aceitos quando há mudança de composição ao longo do tempo? Um CEG raiz é uma chave operacional válida?')
bullet('Qual é a vigência correta de capacidade, localização e associação espacial para cada hora histórica?')
bullet('Quantos meses e anos são o mínimo para representar o ciclo anual e os regimes relevantes? Há períodos oficiais para reservar como teste final?')
bullet('O modelo deve generalizar para conjuntos não vistos, ou somente para conjuntos com histórico e cadastro aprovados?')
doc.add_heading('Operação e engenharia elétrica', level=2)
bullet('Quais métricas e erros por conjunto seriam aceitáveis para um estudo de fluxo de potência? Um ganho estatístico pequeno muda alguma decisão no ANAREDE?')
bullet('Como o especialista prefere receber incerteza: intervalo por conjunto, faixa de MW, aviso fora do domínio ou cenários baixo/base/alto?')
bullet('Quais campos do PWF podem mudar no estudo? O contrato atual preserva tamanho, codificação, DGER e DGEI e altera somente Pg; isso atende ao procedimento do mentor?')
bullet('Qual caso ANAREDE não sigiloso pode ser usado para o aceite e quais condições de convergência devem ser verificadas fora do ClimaGrid?')
doc.add_heading('Decisão de adoção', level=2)
bullet('É preferível um modelo híbrido residual global, modelos por cluster ou modelos por conjunto? Como equilibrar desempenho e manutenção?')
bullet('Qual é o gatilho para sair de shadow mode e qual é o gatilho de rollback para a curva física?')
bullet('Que mudança mínima na estimativa justifica trocar a curva no processo do analista: redução de erro, melhor explicação de gargalos, menor retrabalho ou outro benefício?')

doc.add_heading('10 Roteiro de fala para abrir a mentoria', level=1)
p('“O ClimaGrid já consegue pegar uma hora de cenário climático, estimar potência por conjunto e gerar um PWF que altera somente Pg. O replay histórico é observado e não usa ML. No cenário climático, o caminho servido ainda é uma curva física com clipping por capacidade e disponibilidade. Em paralelo, testamos um LightGBM residual: ele corrige a curva e volta a respeitar os limites físicos. Os testes de agosto de 2024 mostram melhora no recorte experimental, mas o alvo, a disponibilidade e a cobertura histórica ainda não foram aprovados para produção. Quero validar primeiro qual grandeza devemos estimar e qual evidência faria o modelo ser aplicável no ANAREDE. Minha hipótese é manter a curva como guardrail e fallback, e promover o ML apenas se ele demonstrar ganho estável e impacto operacional.”')

doc.add_heading('11 Fonte e classificação das evidências', level=1)
p('Contexto canônico e estado atual: Docs/CONTEXTO_PROJETO_IA.md, AGENTS.md, README.md, backend/README.md e backend/ai-service/README.md. Esses documentos prevalecem sobre a visão antiga quando há conflito.')
p('Implementação do cenário e do estimador: backend/ai-service/app/climate_file.py, backend/ai-service/app/predictor.py, backend/ai-service/training/physical_curve.py, backend/ai-service/training/train.py e backend/ai-service/training/build_dataset.py.')
p('Experimentos: Docs/ML/EXPERIMENTO_02_USINA_01.md, Docs/ML/EXPERIMENTOS_03_A_06_MULTIUSINA.md, Docs/ML/AUDITORIA_FASE_2.md e Docs/ML/PLANO_PRATICO_DADOS_E_TREINO_FASE_2.md, além dos artefatos locais em backend/ai-service/artifacts/experiments. Os artefatos são evidência de experimento, não aprovação de produção.')
p('Visão original do Ideathon: Docs/ClimaGrid_Documento_Tecnico.pdf e Docs/Equipe 16 - Entrega Final Ideathon ClimaGrid.pdf. Esses PDFs foram conferidos como material de proposta; a leitura do estado atual segue o contexto canônico.')

doc.core_properties.title = 'ClimaGrid e o estimador vento para potência'
doc.core_properties.subject = 'Briefing de mentoria técnica sobre ML e viabilidade do projeto'
doc.core_properties.author = 'Equipe ClimaGrid'
doc.core_properties.comments = 'Documento de preparação para mentoria técnica'
doc.save(OUT)
print(OUT.resolve())
